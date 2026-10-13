"""docx에 본문 무수정으로 Word 댓글을 단다.

사용:
    uv run --quiet --with 'python-docx>=1.2' python -I add_comments.py spec.json

spec.json:
{
  "author": "홍길동",
  "initials": "HGD",
  "in_dir": "in",
  "out_dir": "out",
  "suffix": "_검토",
  "files": [
    {"glob": "draft*.docx",
     "comments": [
       {"anchor": "[발명자 초안검토 의견]", "text": "검토 의견입니다.\n..."},
       {"anchor": "【청구항 2】", "text": "...", "nth": 0},
       {"anchor": "[수식 1]", "anchor_end": "여기서 Tk는", "text": "여러 문단을 한 댓글 범위로"}
     ]}
  ]
}

anchor는 문단 텍스트에 포함된 문자열이고 표 셀도 찾는다. nth는 같은 anchor가 여럿일 때 순번.
anchor_end(와 nth_end)를 주면 anchor 문단부터 anchor_end 문단까지 한 댓글 범위로 묶는다.
빨간 글씨가 여러 문단에 걸치면 이렇게 묶어야 모든 빨간 문단이 댓글 범위 안에 든다.
같은 문구가 명세서 본문과 청구항에 둘 다 있으면 nth로 청구항 쪽을 고른다(발송 전 red 문단 커버리지 확인).
text 안의 줄바꿈은 댓글 안 문단 구분이 된다.

저장 뒤 자동 검증: 본문 문단과 표 셀 텍스트가 원본과 동일한지, 댓글 수, 금지 문자(em dash,
en dash, 가운뎃점, 화살표), 내부 줄 번호 참조([155], [0113] 꼴), 분류 라벨과 판정어.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document

FORBIDDEN = "—–·→←⇒"
LINE_REF = re.compile(r"\[0?\d{3}\]")
TONE = re.compile(r"(^|\n)(답변\.|권고\.)|맞습니다|적절하|맞아 보입니다")


def iter_table(t, seen):
    for row in t.rows:
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            yield from cell.paragraphs
            for nt in cell.tables:
                yield from iter_table(nt, seen)


def iter_paragraphs(doc):
    yield from doc.paragraphs
    seen = set()
    for t in doc.tables:
        yield from iter_table(t, seen)


def body_text(doc):
    return [p.text for p in iter_paragraphs(doc)]


def find(doc, anchor, nth):
    hits = [p for p in iter_paragraphs(doc) if anchor in p.text]
    return hits[nth] if len(hits) > nth else None


def add(doc, author, initials, anchor, text, nth=0, anchor_end=None, nth_end=0):
    p = find(doc, anchor, nth)
    if p is None:
        return False, f"MISSING anchor: {anchor!r}"
    runs = [r for r in p.runs if r.text]
    if not runs:
        return False, f"NO RUNS: {anchor!r}"
    last = runs[-1]
    if anchor_end is not None:
        q = find(doc, anchor_end, nth_end)
        if q is None:
            return False, f"MISSING anchor_end: {anchor_end!r}"
        qruns = [r for r in q.runs if r.text]
        if not qruns:
            return False, f"NO RUNS: {anchor_end!r}"
        last = qruns[-1]
    doc.add_comment(runs=(runs[0], last), text=text, author=author, initials=initials)
    return True, f"ok: {anchor[:30]!r}"


def verify(src, dst, expected):
    a, b = Document(str(src)), Document(str(dst))
    same = body_text(a) == body_text(b)
    bodies = [c.text for c in b.comments]
    joined = "\n".join(bodies)
    problems = []
    if not same:
        problems.append("BODY CHANGED")
    if len(bodies) != expected:
        problems.append(f"COMMENT COUNT {len(bodies)} != {expected}")
    bad = [c for c in joined if c in FORBIDDEN]
    if bad:
        problems.append(f"FORBIDDEN CHARS {len(bad)}")
    refs = LINE_REF.findall(joined)
    if refs:
        problems.append(f"LINE REFS {refs}")
    tone = TONE.findall(joined)
    if tone:
        problems.append(f"TONE HITS {len(tone)}")
    return problems


def main(spec_path):
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    in_dir, out_dir = Path(spec["in_dir"]), Path(spec["out_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    author, initials = spec["author"], spec.get("initials", "")
    suffix = spec.get("suffix", "_검토")
    failed = False
    for f in spec["files"]:
        src = next(in_dir.glob(f["glob"]))
        doc = Document(str(src))
        print(f"== {src.name}")
        n = 0
        for c in f["comments"]:
            ok, msg = add(doc, author, initials, c["anchor"], c["text"], c.get("nth", 0),
                          c.get("anchor_end"), c.get("nth_end", 0))
            n += ok
            print("  ", msg)
            failed |= not ok
        dst = out_dir / src.name.replace(".docx", f"{suffix}.docx")
        doc.save(str(dst))
        problems = verify(src, dst, n)
        print("  saved", dst.name, "comments:", n, "| verify:", problems or "ok")
        failed |= bool(problems)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main(sys.argv[1])
